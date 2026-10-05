import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import zipfile, os, shutil, re

docx_path = 'Laporan_UTS_Blockchain_237006079.docx'

# Step 1: Update image2.png in docx zip with diagram3_merkle_tree.png
extract_dir = 'docx_extract_temp'
if os.path.exists(extract_dir):
    shutil.rmtree(extract_dir)

with zipfile.ZipFile(docx_path, 'r') as z:
    z.extractall(extract_dir)

shutil.copyfile('diagram3_merkle_tree.png', os.path.join(extract_dir, 'word/media/image2.png'))

with zipfile.ZipFile(docx_path, 'w', zipfile.ZIP_DEFLATED) as z:
    for root, dirs, files in os.walk(extract_dir):
        for file in files:
            full_path = os.path.join(root, file)
            rel_path = os.path.relpath(full_path, extract_dir)
            z.write(full_path, rel_path)

shutil.rmtree(extract_dir)
print('Image2 (Diagram 3) updated in docx zip.')

# Step 2: Load document and apply document structure updates
doc = docx.Document(docx_path)

# Helper function to format cell
def set_cell(cell, text, font_size_pt=7.0, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, v_align=WD_ALIGN_VERTICAL.TOP):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    run.font.italic = italic
    cell.vertical_alignment = v_align

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='000000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = parse_xml(f'<{tag} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="0" w:color="{edge_data.get("color", "000000")}"/>')
            tcBorders.append(element)
    tcPr.append(tcBorders)

# ==========================================
# REVISI 1: TABEL 2 (Catatan Pelaksanaan)
# ==========================================
t2 = doc.tables[2]

# Delete P21 (long introductory paragraph before Table 2)
# Find P21:
for p in doc.paragraphs:
    if p.text.startswith('Pelaksanaan perancangan, pengembangan prototipe'):
        p.text = '' # Clear text so it doesn't take space
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        break

# Configure Table 2: 3 columns, 7 rows (1 header + 6 data)
# While Table 2 currently has 9 rows, delete rows 7 and 8
while len(t2.rows) > 7:
    tr = t2.rows[-1]._tr
    t2._tbl.remove(tr)

# Set Header Row
headers_t2 = ['Tahap', 'Kegiatan yang Dilakukan', 'Hasil dan Rujukan Bukti']
for i, h in enumerate(headers_t2):
    set_cell(t2.rows[0].cells[i], h, font_size_pt=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# Repeating header & cantSplit
header_trPr = t2.rows[0]._tr.get_or_add_trPr()
header_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} />'))

for r in t2.rows:
    trPr = r._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} />'))

t2_data = [
    (
        'Tahap 1: Analisis dan Desain',
        'Identifikasi 4 pemangku kepentingan ekosistem ijazah, evaluasi kelemahan proses verifikasi manual konvensional, pemodelan 6 ancaman STRIDE, serta perumusan arsitektur jaringan konsorsium permissioned (EVM-compatible / Hyperledger Besu) dan skema Zero-PII on-chain.',
        'Dokumen design-decisions.md, Gambar 1 (Arsitektur Jaringan Konsorsium), Tabel 3 (Komparasi DB vs Blockchain), dan Tabel 10 (Matriks Ancaman STRIDE).'
    ),
    (
        'Tahap 2: Lingkungan dan Kontrak',
        'Inisialisasi repositori Hardhat lokal (Chain ID 31337), implementasi smart contract AcademicCredentialRegistry.sol (OpenZeppelin v5 AccessControl, Pausable, MerkleProof), kompilasi bytecode Solidity 0.8.28, dan deployment kontrak.',
        'Kontrak aktif di alamat 0x5FbDB2315678afecb367f032d93F642f64180aa3, Tx: 0xea3e596f... Gas: 1.039.368 unit [B-01]. (Eksekusi: 2 Oktober 2026, 18:02:25 WIB, demo-output.txt).'
    ),
    (
        'Tahap 3: Data dan Kriptografi',
        'Penyusunan dataset 6 lulusan sintetis, kanonikalisasi JSON mengacu pada RFC 8785, injeksi salt 256-bit acak CSPRNG (crypto.randomBytes), kalkulasi komitmen daun Keccak-256, dan pembentukan Merkle tree off-chain.',
        '6 Daun bergaram terhitung, Merkle Root: 0x6879c11e179813273132ad5cb23b58bd713d41c6535c41fec7f54ca4f3a70647 [B-02], Tabel 5, Gambar 2. (Eksekusi: 2 Oktober 2026, 18:02:25 WIB, demo-output.txt).'
    ),
    (
        'Tahap 4: Eksekusi dan Pengujian',
        'Eksekusi penjangkaran batch wisuda on-chain oleh ISSUER_ROLE, verifikasi ijazah sah via eth_call, uji ketahanan tampering klaim IPK, uji otorisasi RBAC penyerang, pencabutan kredensial oleh REVOKER_ROLE, dan audit jeda darurat (pause/unpause).',
        'Root terdaftar Block #2 tx 0xaf555924... [B-03]; verifikasi valid 0 gas [B-04]; manipulasi ditolak [B-05A]; peretas dicegat [B-05B]; pencabutan [B-06]; pause [B-07]. 8 dari 8 unit test lulus (T-01 s.d. T-08). Log: demo-output.txt (2 Okt 2026 18:02:25 WIB) dan test-output.txt (2 Okt 2026 18:01:47 WIB).'
    ),
    (
        'Tahap 5: Pengukuran Performa',
        'Automated benchmarking 30 iterasi terukur (+ 2 warm-up) pada batch skala N=6, 100, dan 1.000 lulusan sintetis: waktu build pohon off-chain, gas issueBatch, gas revoke, biaya calldata EIP-2028, dan estimasi gas eksekusi murni verifikasi.',
        'Gas eksekusi murni issueBatch konstan 164.911 gas (O(1)), bukti selisih calldata deterministik tepat 84 gas, model anggaran verifikasi linier [B-08], Tabel 9, Tabel 11, Tabel 13. Log: perf-output.txt (Eksekusi: 2 Oktober 2026, 19:11:56 WIB).'
    ),
    (
        'Tahap 6: Penyusunan Laporan',
        'Penyusunan laporan teknis pada template resmi, perapian rujukan akademik, sinkronisasi metrik B-01 s.d. B-08, render diagram monokrom Graphviz 300 DPI, penandatanganan orisinalitas, ekspor PDF, dan verifikasi visual menyeluruh.',
        'Dokumen laporan teknis lengkap Laporan_UTS_Blockchain_237006079.docx dan berkas PDF final (4 Oktober 2026).'
    )
]

# Column widths: 3.2 cm, 7.2 cm, 6.0 cm
col_widths = [Inches(1.26), Inches(2.83), Inches(2.36)]

for row_idx, row_data in enumerate(t2_data, start=1):
    for col_idx, cell_value in enumerate(row_data):
        cell = t2.rows[row_idx].cells[col_idx]
        set_cell(cell, cell_value, font_size_pt=7.0, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

for row in t2.rows:
    for col_idx, width in enumerate(col_widths):
        row.cells[col_idx].width = width
        # Set thin black border
        set_cell_border(row.cells[col_idx], top=dict(sz=4, val='single', color='000000'),
                                           bottom=dict(sz=4, val='single', color='000000'),
                                           left=dict(sz=4, val='single', color='000000'),
                                           right=dict(sz=4, val='single', color='000000'))

# Insert paragraphs below Table 2
# Find element after Table 2 in document body
t2_elem = t2._element
p_after_t2_1 = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="60" w:after="40" w:line="250" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="17"/><w:szCs w:val="17"/></w:rPr></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="17"/><w:szCs w:val="17"/></w:rPr><w:t>Setiap artefak bukti pelaksanaan diberi kode referensi unik B-01 sampai B-08 dan merujuk langsung ke berkas log keluaran eksekusi nyata pada repositori prototipe (demo-output.txt, test-output.txt, dan perf-output.txt). Seluruh proses perancangan, otomatisasi pengujian, dan penyusunan laporan ini dibantu oleh agen kecerdasan artifisial (AI) sebagaimana diungkapkan secara jujur dan transparan pada Lampiran B.</w:t></w:r></w:p>')

p_after_t2_2 = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="40" w:after="100" w:line="250" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="17"/><w:szCs w:val="17"/></w:rPr></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:b/><w:sz w:val="17"/><w:szCs w:val="17"/></w:rPr><w:t>Aspek Keamanan dan Kerahasiaan Lingkungan Uji: </w:t></w:r><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="17"/><w:szCs w:val="17"/></w:rPr><w:t>Seluruh pengujian menggunakan data kelulusan sintetis fiktif sehingga tidak ada data pribadi riil yang terekspos; pengujian tidak menggunakan kunci privat bernilai nyata, melainkan memanfaatkan akun pengembang bawaan node Hardhat lokal; serta nilai salt 256-bit dibangkitkan secara acak murni menggunakan CSPRNG Node.js (crypto.randomBytes) pada setiap sesi eksekusi guna menjamin sifat unik daun komitmen tanpa kebocoran entropi.</w:t></w:r></w:p>')

t2_elem.addnext(p_after_t2_2)
t2_elem.addnext(p_after_t2_1)

print('Table 2 and post-table notes updated successfully.')

# ==========================================
# REVISI 2: LAMPIRAN A (Pernyataan Orisinalitas) & TABEL TANDA TANGAN
# ==========================================

# Find P111 (Lampiran A) and P112 (statement)
p_lamp_a = None
p_stmt = None
for p in doc.paragraphs:
    if 'Lampiran A: Pernyataan Orisinalitas' in p.text:
        p_lamp_a = p
    elif 'Pernyataan Integritas Akademik:' in p.text or 'Saya menyatakan bahwa jawaban ujian' in p.text:
        p_stmt = p

if p_lamp_a:
    p_lamp_a.paragraph_format.page_break_before = True
    p_lamp_a.paragraph_format.keep_with_next = True
    p_lamp_a.paragraph_format.space_before = Pt(6)
    p_lamp_a.paragraph_format.space_after = Pt(4)

if p_stmt:
    p_stmt.text = '' # Clear
    p_stmt.paragraph_format.keep_with_next = True
    p_stmt.paragraph_format.space_before = Pt(2)
    p_stmt.paragraph_format.space_after = Pt(8)
    p_stmt.paragraph_format.line_spacing = 1.15
    run = p_stmt.add_run('Saya menyatakan bahwa jawaban ujian dan mini-project ini merupakan hasil kerja saya sendiri. Seluruh sumber dan penggunaan alat bantu telah saya ungkapkan secara jujur sesuai ketentuan mata kuliah.')
    run.font.name = 'Calibri'
    run.font.size = Pt(10.0)

# Replace Table 14 with a 4-column signature table
t14 = doc.tables[14]

# Remove old Table 14 from document body
t14_parent = t14._element.getparent()
t14_index = list(t14_parent).index(t14._element)

# Create a new table directly in its place
new_tbl = doc.add_table(rows=2, cols=4)
new_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_sig = ['Nama', 'NIM', 'Tanggal', 'Tanda Tangan']
sig_col_widths = [Inches(1.77), Inches(1.18), Inches(1.38), Inches(2.13)] # 4.5cm, 3.0cm, 3.5cm, 5.4cm

for i, h in enumerate(headers_sig):
    cell = new_tbl.rows[0].cells[i]
    set_cell(cell, h, font_size_pt=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)
    set_cell_border(cell, top=dict(sz=6, val='single', color='000000'),
                          bottom=dict(sz=6, val='single', color='000000'),
                          left=dict(sz=6, val='single', color='000000'),
                          right=dict(sz=6, val='single', color='000000'))

# Data row
data_sig = ['Fajar Geran Arifin', '237006079', '4 Oktober 2026', '']
for i, d in enumerate(data_sig[:3]):
    cell = new_tbl.rows[1].cells[i]
    set_cell(cell, d, font_size_pt=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, v_align=WD_ALIGN_VERTICAL.CENTER)
    set_cell_border(cell, top=dict(sz=6, val='single', color='000000'),
                          bottom=dict(sz=6, val='single', color='000000'),
                          left=dict(sz=6, val='single', color='000000'),
                          right=dict(sz=6, val='single', color='000000'))

# Cell 3: Tanda Tangan
cell_ttd = new_tbl.rows[1].cells[3]
cell_ttd.text = ''
cell_ttd.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p_ttd = cell_ttd.paragraphs[0]
p_ttd.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ttd.paragraph_format.space_before = Pt(2)
p_ttd.paragraph_format.space_after = Pt(2)
p_ttd.paragraph_format.line_spacing = 1.0
run_ttd = p_ttd.add_run()
run_ttd.add_picture('ttd_penulis.png', width=Cm(3.8))
set_cell_border(cell_ttd, top=dict(sz=6, val='single', color='000000'),
                          bottom=dict(sz=6, val='single', color='000000'),
                          left=dict(sz=6, val='single', color='000000'),
                          right=dict(sz=6, val='single', color='000000'))

# Set widths and cantSplit
for r in new_tbl.rows:
    trPr = r._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} />'))
    for i, w in enumerate(sig_col_widths):
        r.cells[i].width = w

# Set row 1 height to 4.2 cm so signature is roomy
trPr1 = new_tbl.rows[1]._tr.get_or_add_trPr()
trPr1.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="2400" w:hRule="atLeast"/>'))

# Move new_tbl element to where old Table 14 was
t14_parent.insert(t14_index, new_tbl._element)
t14_parent.remove(t14._element)

print('Table 14 replaced with 4-column signature table.')

# ==========================================
# CEK KONSISTENSI HASHES & RFC 8785 DI DOKUMEN
# ==========================================

# Table 5: Leaf Hashes
t5 = doc.tables[5]
t5_leaves = [
    ('0xc4d1a22493f0f84e... (32 byte)', 'Valid Terverifikasi (Bukti B-04)'),
    ('0x7f0dc81a7a1ee238... (32 byte)', 'Tercabut Resmi (Bukti B-06)'),
    ('0x067083cc811ecdac... (32 byte)', 'Valid dalam Root Batch'),
    ('0xeed86c98124312bf... (32 byte)', 'Valid dalam Root Batch'),
    ('0x1d8e13d64338026c... (32 byte)', 'Valid dalam Root Batch'),
    ('0xffb33bd42ca39de6... (32 byte)', 'Valid dalam Root Batch')
]
for idx, (leaf_hash, status) in enumerate(t5_leaves, start=1):
    t5.rows[idx].cells[3].text = leaf_hash
    p = t5.rows[idx].cells[3].paragraphs[0]
    p.runs[0].font.name = 'Calibri'
    p.runs[0].font.size = Pt(7.0)

# Table 8: Transaction Hash
t8 = doc.tables[8]
t8.rows[3].cells[1].text = 'Gas Limit: 30.000.000; Gas Used: 211.084 gas (penjangkaran batch on-chain) / 187.873 s.d. 187.957 gas (benchmark); Status: 1 (Success).'
p8_3 = t8.rows[3].cells[1].paragraphs[0]
p8_3.runs[0].font.name = 'Calibri'
p8_3.runs[0].font.size = Pt(7.5)

t8.rows[4].cells[1].text = 'Event: BatchIssued(WISUDA-2026-PERIODE-1, 0x6879c11e..., timestamp); Tx Hash: 0xaf555924fd6db1366a269faabd6dcc2efedce4a7177a9674c8bee194bf777dc9 [B-03].'
p8_4 = t8.rows[4].cells[1].paragraphs[0]
p8_4.runs[0].font.name = 'Calibri'
p8_4.runs[0].font.size = Pt(7.5)

# Text updates: RFC 8785 phrasing
for p in doc.paragraphs:
    if 'kanonikalisasi JSON RFC 8785' in p.text:
        for r in p.runs:
            r.text = r.text.replace('kanonikalisasi JSON RFC 8785', 'kanonikalisasi JSON yang mengacu pada RFC 8785')
    if 'Kanonikalisasi JSON berdasarkan RFC 8785' in p.text:
        for r in p.runs:
            r.text = r.text.replace('Kanonikalisasi JSON berdasarkan RFC 8785', 'Kanonikalisasi JSON yang mengacu pada RFC 8785')
    if 'RFC8785_Canonicalize' in p.text:
        for r in p.runs:
            r.text = r.text.replace('RFC8785_Canonicalize', 'Canonicalize_Acuan_RFC8785')

for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            if 'modul kanonikalisasi RFC 8785' in c.text:
                for p in c.paragraphs:
                    for r_run in p.runs:
                        r_run.text = r_run.text.replace('modul kanonikalisasi RFC 8785', 'modul kanonikalisasi mengacu pada RFC 8785')

# Save updated docx
doc.save(docx_path)
print('Document updated and saved successfully.')
