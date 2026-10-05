import docx

doc = docx.Document('uts-blockchain-submission/Laporan_UTS_Blockchain_237006079_v3.docx')

def dump_table(t_idx):
    tbl = doc.tables[t_idx]
    print(f"\n==================== TABLE {t_idx} ({len(tbl.rows)} rows, {len(tbl.columns)} cols) ====================")
    for r_idx, row in enumerate(tbl.rows):
        row_str = []
        for c_idx, cell in enumerate(row.cells):
            txt = cell.text.strip().replace("\n", " ")
            p = cell.paragraphs[0] if cell.paragraphs else None
            font_name = p.runs[0].font.name if (p and p.runs and p.runs[0].font.name) else "default"
            font_size = p.runs[0].font.size.pt if (p and p.runs and p.runs[0].font.size) else "default"
            bold = p.runs[0].font.bold if (p and p.runs) else False
            row_str.append(f"C{c_idx}[{font_name},{font_size}pt,b={bold}]: {txt[:50]}")
        print(f"R{r_idx}: " + " | ".join(row_str))

for t in [2, 5, 8, 9, 11, 12, 13]:
    dump_table(t)

