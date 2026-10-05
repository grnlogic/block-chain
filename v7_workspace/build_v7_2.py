#!/usr/bin/env python3
"""
build_v7_2.py
Master build script to generate Laporan_UTS_Blockchain_237006079_v7.2.docx and .pdf
directly and reproducibly from the base docx:
uts-blockchain-submission/Laporan_UTS_Blockchain_237006079.docx.

Contains ALL changes (v7, v7.1, v7.2) end-to-end without manual docx editing:
- v7: Table 11/13 move to Lampiran D/E, renumbering body tables 1-10, UTXO compression,
      Diagram 2 scaling (80%), soft-pedaling claims.
- v7.1: BlockGasLimit 60.000.000 (0x3938700), hardware specs 12th Gen Intel(R) Core(TM)
        i5-12450H 12 thread logis, RAM total 15 GiB, Table 2/7 ellipsis formatting,
        UU PDP separate articles (Pasal 4 ayat 3, Pasal 8, Pasal 9, Pasal 43 ayat 1).
- v7.2: Body strictly <= 12 pages (Cara i), 11 blank lines remaining on page 12,
        Lampiran F for B-06 (exp_b06_output.txt & Run A-E table), Lampiran D & E raised to 9 pt,
        body tables untouched, short hex notation restored (0x03e8, 0x06, 0x3938700),
        Lampiran A AI disclosure updated, PDF metadata set.
"""

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import subprocess
import os
import re

SRC_DOCX = "uts-blockchain-submission/Laporan_UTS_Blockchain_237006079.docx"
OUT_DOCX = "v7_workspace/Laporan_UTS_Blockchain_237006079_v7.2.docx"
OUT_PDF = "v7_workspace/Laporan_UTS_Blockchain_237006079_v7.2.pdf"

print(f"Loading base docx from: {SRC_DOCX}")
doc = docx.Document(SRC_DOCX)

# Set Document Core Properties (Metadata)
doc.core_properties.author = "Fajar Geran Arifin"
doc.core_properties.title = "UTS Blockchain - Rancangan Sistem Verifikasi Kredensial Akademik"

def replace_in_p(p, old, new, preserve_format=True):
    """Replace text in paragraph preserving runs or formatting."""
    if old not in p.text:
        return False
    if not p.runs:
        p.text = p.text.replace(old, new)
        return True
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    # If old spans across runs
    fname = p.runs[0].font.name or "Calibri"
    fsize = p.runs[0].font.size or Pt(10.0)
    color = p.runs[0].font.color.rgb if p.runs[0].font.color else None
    bold = p.runs[0].font.bold
    p.text = p.text.replace(old, new)
    if preserve_format:
        for r in p.runs:
            r.font.name = fname
            r.font.size = fsize
            if color:
                r.font.color.rgb = color
            r.font.bold = bold
    return True

# =============================================================================
# 1. FOOTER UPDATE
# =============================================================================
for sec in doc.sections:
    footer = sec.footer
    for p in footer.paragraphs:
        if "Template Laporan Kegiatan Blockchain" in p.text:
            p.text = "Laporan UTS Blockchain | TA 2026/2027"

# =============================================================================
# 2. CLEAR CANTSPLIT ON ALL TABLE ROWS FOR NATURAL PAGE FLOW
# =============================================================================
for t in doc.tables:
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr()
        cs = trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit")
        if cs is not None:
            trPr.remove(cs)

# =============================================================================
# 3. HEADING FIXES (REMOVE "Case n")
# =============================================================================
heading_fixes = {
    "4.1 Analisis Kelayakan, Pemangku Kepentingan, dan Arsitektur Jaringan (P1, P2 / Case 1)":
        "4.1 Analisis Kelayakan, Pemangku Kepentingan, dan Arsitektur Jaringan (P1, P2)",
    "4.2 Kriptografi, W3C Verifiable Credentials, dan Integritas Daun Bergaram (P3 / Case 2)":
        "4.2 Kriptografi, W3C Verifiable Credentials, dan Integritas Daun Bergaram (P3)",
    "4.3 Jaringan P2P, Siklus Transaksi, dan Mekanisme Konsensus (P2, P3 / Case 3)":
        "4.3 Jaringan P2P, Siklus Transaksi, Mekanisme Konsensus, dan Evaluasi Model UTXO (P2, P3)",
    "4.5 Model Akun EVM, Analisis Ketat Komponen Gas, dan Model Regresi (P3, P5 / Case 5)":
        "4.4 Model Akun EVM, Analisis Ketat Komponen Gas, dan Model Regresi (P3, P5)",
    "4.6 Analisis Keamanan, Model Ancaman STRIDE, dan Respons Insiden 24 Jam (P4)":
        "4.5 Analisis Keamanan, Model Ancaman STRIDE, dan Respons Insiden 24 Jam (P4)",
    "4.7 Implementasi Prototipe, Kriteria Penerimaan, dan Pengujian Performa (P5)":
        "4.6 Implementasi Prototipe, Kriteria Penerimaan, dan Pengujian Performa (P5)",
    "4.8 Keputusan Desain, Batasan Teknis, Pertimbangan Etika, dan Refleksi (P6)":
        "4.7 Keputusan Desain, Batasan Teknis, Pertimbangan Etika, dan Refleksi (P6)"
}
for p in doc.paragraphs:
    t = p.text.strip()
    if t in heading_fixes:
        p.text = heading_fixes[t]

# =============================================================================
# 4. COMPRESS SUBCHAPTER 4.4 (UTXO) INTO 4.3 AND REMOVE TABLE 7
# =============================================================================
p_utxo_heading = None
p_utxo_p1 = None
p_utxo_p2 = None

for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("4.4 Evaluasi Model UTXO Bitcoin"):
        p_utxo_heading = p
        if i + 1 < len(doc.paragraphs):
            p_utxo_p1 = doc.paragraphs[i+1]
        if i + 2 < len(doc.paragraphs):
            p_utxo_p2 = doc.paragraphs[i+2]
        break

for p in doc.paragraphs:
    if "Analisis Fork & Pemilihan Konsensus:" in p.text:
        p.text = p.text + (
            " Evaluasi Model UTXO vs Akun EVM (P2): Dalam memenuhi luaran P2 terkait justifikasi arsitektur, "
            "model Unspent Transaction Output (UTXO) Bitcoin dianalisis sebagai pembanding. Model UTXO bersifat stateless "
            "dan membatasi penyisipan data pada script OP_RETURN maksimal 80 byte (~0.005 BTC biaya transaksi ilustratif tanpa smart logic). "
            "Kelemahan fatal UTXO untuk ijazah adalah ketiadaan penyimpanan status dinamis (stateful storage); untuk membatalkan ijazah pada UTXO, "
            "penerbit harus membuat rantai transaksi pembelanjaan baru yang rumit. Sebaliknya, model akun EVM menyediakan mesin status "
            "terprogram yang memungkinkan penyimpanan status pencabutan individual secara langsung dalam slot mapping storage (_revokedCredentials[leafHash]) "
            "dengan kompleksitas akses O(1) dan biaya komputasi deterministik."
        )

if p_utxo_heading is not None:
    p_utxo_heading._element.getparent().remove(p_utxo_heading._element)
if p_utxo_p1 is not None and "Analisis Model UTXO:" in p_utxo_p1.text:
    p_utxo_p1._element.getparent().remove(p_utxo_p1._element)
if p_utxo_p2 is not None and "Komparasi Model Akun vs UTXO:" in p_utxo_p2.text:
    p_utxo_p2._element.getparent().remove(p_utxo_p2._element)

# Remove Table 7 (UTXO evaluation table)
tbl7_utxo = doc.tables[7]
tbl7_utxo._element.getparent().remove(tbl7_utxo._element)

# =============================================================================
# 5. MOVE TABLE 11 (CRITERIA) & TABLE 13 (BENCHMARK) TO APPENDICES (LAMPIRAN D & E)
# =============================================================================
tbl_criteria = None
tbl_unittest = None
tbl_bench = None

for t in doc.tables:
    first_cell = t.rows[0].cells[0].text.strip()
    if "Kelompok Uji" in first_cell:
        tbl_criteria = t
    elif "ID Tes" in first_cell:
        tbl_unittest = t
    elif "Parameter Pengujian" in first_cell:
        tbl_bench = t

# Move tbl_criteria to Lampiran D
p_lamp_d = doc.add_paragraph()
p_lamp_d.text = "Lampiran D: Matriks Evaluasi Lengkap Kriteria Penerimaan dan Luaran OBE (T-01 s.d. T-08)"
p_lamp_d.paragraph_format.space_before = Pt(14)
p_lamp_d.paragraph_format.space_after = Pt(4)
p_lamp_d.paragraph_format.keep_with_next = True
for r in p_lamp_d.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

tbl_criteria._element.getparent().remove(tbl_criteria._element)
doc.element.body.append(p_lamp_d._element)
doc.element.body.append(tbl_criteria._element)

# Move tbl_bench to Lampiran E
p_lamp_e = doc.add_paragraph()
p_lamp_e.text = "Lampiran E: Rekapitulasi Lengkap Metrik Pengujian Performa 30 Run Terukur"
p_lamp_e.paragraph_format.space_before = Pt(14)
p_lamp_e.paragraph_format.space_after = Pt(4)
p_lamp_e.paragraph_format.keep_with_next = True
for r in p_lamp_e.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

tbl_bench._element.getparent().remove(tbl_bench._element)
doc.element.body.append(p_lamp_e._element)
doc.element.body.append(tbl_bench._element)

# =============================================================================
# 6. HARDWARE SPECS (12th Gen Intel(R) Core(TM) i5-12450H, 12 thread logis, RAM total 15 GiB)
# =============================================================================
# Table 1 update
t1 = doc.tables[1]
for r in t1.rows:
    for c in r.cells:
        if "Intel" in c.text and "RAM" in c.text:
            for p in c.paragraphs:
                for run in p.runs:
                    run.text = "Perangkat keras: 12th Gen Intel(R) Core(TM) i5-12450H (12 thread logis), RAM total 15 GiB. OS: Ubuntu 24.04.4 LTS x86_64, Node.js v24.14.0."

# =============================================================================
# 7. TABLE 2 UPDATES (Address & Merkle Root ellipsis & B-01, B-06, B-07 numerics)
# =============================================================================
t2 = doc.tables[2]
t2.rows[2].cells[2].text = "Kontrak aktif di alamat 0x5FbDB23156..., Tx: 0x145b5523..., Gas: 1.039.368 unit [B-01]"
for p in t2.rows[2].cells[2].paragraphs:
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)

t2.rows[3].cells[2].text = "Pohon Merkle (129,44 ms off-chain, B-02); 6 daun bergaram, Root: 0x46ca2e83... (Tabel 5, Diagram 3, results/demo-output.txt)"
for p in t2.rows[3].cells[2].paragraphs:
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)

t2.rows[4].cells[2].text = (
    "Root terdaftar Block #2 tx 0x1737c359... [B-03]; verifikasi valid 0 gas [B-04]; "
    "manipulasi ditolak [B-05A/B]; cabut ijazah (Gas: 56.359 unit) [B-06]; "
    "jeda darurat pause (Gas: 48.214 unit) [B-07], Tabel 7, Tabel 9, Tabel 10"
)
for p in t2.rows[4].cells[2].paragraphs:
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)

t2.rows[5].cells[2].text = (
    "Gas eksekusi murni issueBatch konstan 164.911 gas (O(1)), bukti selisih calldata deterministik tepat 84 gas, "
    "model kelayakan verifikasi linear 7.800 + 251*k gas, kueri publik 0 gas. "
    "Rujukan: B-08, Tabel 8, Lampiran D, Lampiran E (results/perf-output.txt)"
)
for p in t2.rows[5].cells[2].paragraphs:
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)

# =============================================================================
# 8. TABLE 7 (FIELD TRANSAKSI EVM): BLOCK GAS LIMIT 60.000.000 & TX HASH ELLIPSIS
# =============================================================================
for t in doc.tables:
    if "Field Transaksi EVM" in t.rows[0].cells[0].text:
        t.rows[3].cells[0].text = "Gas Limit / Gas Used / Status"
        t.rows[3].cells[1].text = "Gas Limit Blok: 60.000.000 (Parameter terukur node lokal Hardhat); Gas Used: 211.084 gas (penjangkaran batch); Status: 1 (Success)"
        t.rows[4].cells[0].text = "Event Logs & Transaction Hash"
        t.rows[4].cells[1].text = "Event: BatchIssued(WISUDA-2026-PERIODE-1, 0x46ca2e83..., timestamp); Tx Hash: 0x1737c359...; Block: #2 [B-03]"
        for r_idx in [3, 4]:
            for c_idx in [0, 1]:
                for p in t.rows[r_idx].cells[c_idx].paragraphs:
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.font.size = Pt(7.5)

# =============================================================================
# 9. GENERAL TEXT UPDATES & CLAIM ADJUSTMENTS
# =============================================================================
for p in doc.paragraphs:
    # Reproducibility intro
    if "Untuk menjamin bahwa seluruh hasil observasi" in p.text or "seluruh konfigurasi perangkat keras" in p.text:
        p.text = (
            "Untuk memastikan bahwa seluruh hasil observasi, angka konsumsi gas, dan waktu komputasi yang dilaporkan "
            "dalam dokumen ini dapat diuji ulang secara independen oleh dosen penguji maupun peneliti lain, seluruh "
            "konfigurasi perangkat keras, perangkat lunak, pustaka dependensi, dan perintah eksekusi dirangkum pada Tabel 1:"
        )

    # Catatan Pelaksanaan intro
    if p.text.strip() == "3 Catatan Pelaksanaan":
        p.text = "3 Catatan Pelaksanaan\nPelaksanaan mini-project perancangan dan evaluasi sistem verifikasi kredensial akademik berbasis blockchain konsorsium ini dirangkum dalam Tabel 2."

    # Section 4.1 Table 3 intro
    if "melainkan dijamin oleh konsensus kriptografis terdistribusi." in p.text:
        p.text = p.text.replace(
            "melainkan dijamin oleh konsensus kriptografis terdistribusi.",
            "melainkan dijamin oleh konsensus kriptografis terdistribusi, sebagaimana dikomparasikan dalam Tabel 3."
        )

    # Section 4.1 Table 4 intro & Citation [7]
    if "seperti diilustrasikan pada Diagram 1." in p.text and "Tabel 4" not in p.text:
        p.text = p.text.replace(
            "seperti diilustrasikan pada Diagram 1.",
            "seperti diilustrasikan pada Diagram 1. Pembagian peran dan penempatan data dirangkum pada Tabel 4."
        )
    if "Membutuhkan layanan verifikasi yang instan (< 1 detik)" in p.text and "[7]" not in p.text:
        p.text = p.text.replace(
            "Membutuhkan layanan verifikasi yang instan (< 1 detik)",
            "Membutuhkan layanan verifikasi yang instan (< 1 detik sesuai batas toleransi respons kognitif Nielsen [7])"
        )

    # Section 4.2 W3C VC, Salt, & UU PDP (Pasal 4 ayat 3, Pasal 8, Pasal 9, Pasal 43 ayat 1)
    if "Standar W3C VC & Kepatuhan Hak Hapus:" in p.text or "Standar W3C VC & Keselarasan Hak Hapus:" in p.text:
        p.text = (
            "Standar W3C VC & Keselarasan Hak Hapus: Kanonikalisasi JSON yang mengacu pada RFC 8785 [9] "
            "memastikan bahwa susunan kunci atribut kamus selalu deterministik secara byte-level di seluruh platform sistem operasi. "
            "Standar W3C Verifiable Credentials Data Model v2.0 [1] mendefinisikan mekanisme credentialStatus untuk memeriksa "
            "keabsahan dokumen; penjangkaran status melalui Merkle Root on-chain [4] merupakan pilihan desain arsitektur pada prototipe ini. "
            "Desain arsitektur Zero-PII ini dirancang selaras dengan prinsip pelindungan data pribadi pada UU No. 27 Tahun 2022 (UU PDP) [10], "
            "khususnya klasifikasi data pribadi umum pada Pasal 4 ayat (3), hak subjek data untuk meminta pemusnahan/penghapusan data pada Pasal 8, "
            "serta kewajiban pengendali data untuk melakukan penghapusan data pribadi pada Pasal 43 ayat (1): "
            "jika alumni mengajukan hak penghapusan data, pemusnahan dokumen lokal dan nilai salt off-chain dirancang untuk mengurangi keterkaitan data "
            "(data linkability) antara subjek dan rekaman transaksi; status hukum atas residu hash pada buku besar publik tetap memerlukan telaah hukum lebih lanjut. "
            "Struktur data record daun dan injeksi salt ini dirangkum dalam Tabel 5."
        )

    # Section 4.3 Table 6 intro
    if "empat keunggulan operasional utama:" in p.text and "Tabel 6" not in p.text:
        p.text = p.text.replace(
            "empat keunggulan operasional utama:",
            "empat keunggulan operasional utama, sebagaimana dirangkum dalam Tabel 6:"
        )

    # Section 4.4 Table 7 intro
    if "Komposisi field transaksi" in p.text and "Tabel" in p.text:
        p.text = (
            "Keterkaitan dengan P3, P5 & Model Akun EVM: Untuk memenuhi luaran P3 (desain smart contract) dan P5 (analisis performa gas mendalam), "
            "subbab ini membedah model komputasi akun EVM. EVM mengkategorikan akun ke dalam dua entitas: Externally Owned Account (EOA) yang dikendalikan oleh "
            "pasangan kunci privat/publik kurva secp256k1, dan Contract Account yang dikendalikan oleh kode bytecode tersimpan. "
            "Komposisi field transaksi issueBatch dan kueri verifikasi dirangkum dalam Tabel 7."
        )

    # Section 4.4 Table 8 intro
    if "Perhitungan Ketat Komponen Gas issueBatch:" in p.text:
        p.text = (
            "Perhitungan Ketat Komponen Gas issueBatch: Sesuai prinsip eksekusi EVM, konsumsi gas transaksi tidak boleh dilaporkan "
            "secara mentah tanpa pemisahan komponen. Gas eksekusi murni internal kontrak dihitung dengan formula baku: "
            "Gas_Eksekusi = gasUsed (Receipt) - 21.000 (Intrinsic Tx) - Biaya_Calldata. Biaya calldata dihitung langsung dari payload "
            "tx.data aktual sesuai aturan EIP-2028 [11] (16 gas per byte non-nol dan 4 gas per byte nol), sebagaimana dirinci pada Tabel 8:"
        )

    # Section 4.4 Calldata claim & Short Hex restoration (0x03e8, 0x06)
    if "Pembuktian Sumber Selisih Calldata 84 Gas:" in p.text:
        p.text = p.text.replace(
            "Pengukuran empiris membuktikan bahwa selisih gas sebesar 84 gas antara batch N=6 dan N=1.000 sama sekali bukan berasal dari eksekusi smart contract, melainkan murni dari akumulasi tiga parameter calldata ABI transaksi:",
            "Pengukuran empiris menunjukkan bahwa selisih gas sebesar 84 gas antara batch N=6 dan N=1.000 dapat dijelaskan secara deterministik melalui akumulasi tiga parameter calldata ABI transaksi:"
        )

    # Section 4.5 Table 9 intro
    if "sebagaimana dirangkum dalam Matriks Ancaman STRIDE." in p.text or "Matriks Analisis Ancaman STRIDE" in p.text:
        p.text = (
            "Keamanan arsitektur dianalisis secara komprehensif menggunakan metodologi STRIDE (Spoofing, Tampering, "
            "Repudiation, Information Disclosure, Denial of Service, Elevation of Privilege) untuk mengidentifikasi potensi ancaman "
            "dan kontrol mitigasi yang dirancang pada smart contract dan sistem off-chain. Kontrol mitigasi menerapkan prinsip hak istimewa terkecil "
            "(Principle of Least Privilege) dari Saltzer dan Schroeder (1975) [5] melalui pemisahan peran berwenang, "
            "sebagaimana dirangkum dalam Tabel 9 (Matriks Analisis Ancaman STRIDE)."
        )

    # Section 4.6 (Implementation & Acceptance Criteria)
    if "sebagaimana dirangkum dalam Tabel Kriteria Penerimaan." in p.text or "sebagaimana dirangkum dalam Tabel 11." in p.text:
        p.text = (
            "Kriteria penerimaan (acceptance criteria) disusun dengan merujuk pada standar ilmiah W3C VC [1], Ethereum JSON-RPC [2], "
            "dan RFC 8785 [9]. Rincian matriks evaluasi lengkap kriteria penerimaan terhadap luaran OBE disajikan pada Lampiran D. "
            "Di lingkungan pengujian otomatis Hardhat, seluruh skenario uji fungsional, keamanan, dan performa (T-01 s.d. T-08) "
            "berhasil lulus 100%, sebagaimana dirangkum pada Tabel 10."
        )

    # Section 4.6 Performance Benchmarks intro
    if "Hasil pengujian performa mendalam melalui 30 run terukur" in p.text:
        p.text = (
            "Pengujian performa mendalam dilakukan melalui 30 run terukur (+ 2 run warm-up) pada skrip scripts/perf-bench.js. "
            "Model regresi kelayakan linier pureVerifyGas(k) = 7.800,04 + 251,26*k gas terkonfirmasi memenuhi batas anggaran kelayakan, "
            "dan kueri verifikasi publik via eth_call mengonsumsi 0 gas dengan latensi rata-rata sub-milidetik (0,82 s.d. 0,89 ms). "
            "Rekapitulasi lengkap metrik statistik benchmark 30 run dirangkum pada Lampiran E."
        )

    # Section 4.7 Hardware specs in paragraph
    if "workstation pengembang (12th Gen Intel" in p.text:
        for r in p.runs:
            if "workstation pengembang" in r.text:
                r.text = re.sub(
                    r"workstation pengembang \([^)]+\)",
                    "workstation pengembang (12th Gen Intel(R) Core(TM) i5-12450H 12 thread logis, RAM total 15 GiB, Ubuntu 24.04.4 LTS, Node.js v24.14.0)",
                    r.text
                )

    # Section 4.7 Trade-off (Remove duplicate percentages 33,3%, 66,7%, <0,36%)
    if "akan menghabiskan lebih dari 20 juta gas untuk 1.000 lulusan" in p.text:
        p.text = (
            "Keputusan Desain & Trade-Off: Tiga keputusan desain fundamental melandasi prototipe ini: "
            "(1) On-chain vs Off-chain: Memilih off-chain document storage dengan on-chain cryptographic commitment (Merkle root). "
            "Pendekatan ini menjamin skalabilitas ekstrem dan kepatuhan penuh terhadap regulasi privasi data. "
            "Penjangkaran individual 1.000 dokumen akan menghabiskan lebih dari 20 juta gas untuk 1.000 lulusan. "
            "Pendekatan Merkle batching mengompresi ribuan lulusan ke dalam 1 Merkle Root 32-byte [4] dengan gas konstan 164.911 gas (hemat > 98%), "
            "memungkinkan jutaan kredensial diterbitkan secara efisien dalam satu transaksi tunggal."
        )

    # Section 4.7 Keterbatasan Teknis title cleanups
    replace_in_p(p, "• 1. Lingkungan Node: Lingkungan Single-Node Hardhat:", "• 1. Lingkungan Single-Node In-Process:")
    replace_in_p(p, "• 1. Lingkungan Node: Lingkungan Single-Node In-Process:", "• 1. Lingkungan Single-Node In-Process:")
    replace_in_p(p, "• 2. Data Uji: Data Sintetis Fiktif:", "• 2. Data Uji Sintetis:")
    replace_in_p(p, "• 3. Perangkat Keras: Karakteristik Mesin Uji Tunggal:", "• 3. Karakteristik Mesin Uji Tunggal:")
    replace_in_p(p, "• 4. Simulasi Gas: Hakikat Nilai estimateGas:", "• 4. Hakikat Nilai estimateGas:")
    replace_in_p(p, "• 5. Slot Storage URI: Alokasi Slot Storage metadataURI:", "• 5. Alokasi Slot Storage metadataURI:")
    replace_in_p(p, "• 5. Slot Storage URI: Alokasi Slot Storage URI:", "• 5. Alokasi Slot Storage metadataURI:")
    replace_in_p(p, "• 6. Batas Gas Blok: Parameter Node Lokal Hardhat vs Hipotesis Konsorsium:", "• 6. Batas Gas Blok Node Lokal vs Konsorsium:")
    replace_in_p(p, "• 7. Pembulatan Calldata: Efek Pembulatan Rata-Rata Calldata (83 vs 84 gas):", "• 7. Efek Pembulatan Calldata (83 vs 84 gas):")

    # Section 4.7 Bullet 6 Gas Limit (Measured 60.000.000 & 0x3938700)
    if "• 6. Batas Gas Blok" in p.text:
        p.text = (
            "• 6. Batas Gas Blok Node Lokal vs Konsorsium: Pengukuran aktual pada node lokal Hardhat melalui RPC "
            "eth_getBlockByNumber menghasilkan blockGasLimit sebesar 60.000.000 gas (0x3938700). Angka 30.000.000 gas "
            "dijadikan contoh hipotetis jaringan konsorsium untuk simulasi kapasitas blok. Penjangkaran individual "
            "1.000 ijazah (~20 juta gas) akan menyerap sepertiga (33,3%) dari batas blok 60 juta gas atau dua pertiga "
            "(66,7%) dari contoh hipotetis 30 juta gas, sedangkan skema Merkle batching pada prototipe ini hanya "
            "memerlukan 211.084 gas (< 0,36% dari batas blok 60 juta gas)."
        )

    # Section 4.7 Bullet 7 Calldata
    if "• 7. Efek Pembulatan Calldata" in p.text:
        p.text = (
            "• 7. Efek Pembulatan Calldata (83 vs 84 gas): Selisih rata-rata calldata ~83,6 gas vs 84 gas transaksi "
            "nominal merupakan efek pembulatan statistik digit nomor iterasi batchId (lihat rincian Subbab 4.4)."
        )

    # Section 4.7 Permendikbudristek & UU PDP
    if "Pertimbangan Etika, Regulasi, & Roadmap:" in p.text:
        p.text = (
            "Pertimbangan Etika, Regulasi, & Roadmap: Implementasi sistem ini menegakkan prinsip keadilan sosial dan non-diskriminasi. "
            "Biaya verifikasi yang ditetapkan 0 gas dirancang agar lulusan dari keluarga prasejahtera dan instansi UMKM/pemberi kerja "
            "lokal dapat memverifikasi keabsahan dokumen tanpa hambatan finansial [2]. Secara regulasi, arsitektur dirancang selaras "
            "dengan Permendikbudristek No. 6 Tahun 2022 (Berita Negara RI 2022 No. 167) terkait keterverifikasian ijazah [13], serta "
            "perlindungan privasi UU No. 27 Tahun 2022 (UU PDP) [10]. Roadmap pengembangan masa depan mencakup integrasi skema "
            "W3C Bitstring Status List v1.0 [8] untuk kompresi pencabutan tingkat lanjut dan penerapan Zero-Knowledge Proofs (BBS+ "
            "Signatures) untuk Selective Disclosure (misal membuktikan 'IPK >= 3.50' tanpa membuka angka IPK eksak)."
        )

    # Section 4.7 Refleksi
    if "Refleksi Individual Penulis (Fajar Geran Arifin" in p.text:
        p.text = (
            "Refleksi Individual Penulis (Fajar Geran Arifin / 237006079): Sebagai mahasiswa penyusun, proses pengerjaan mini-project "
            "ini memberikan pemahaman mendalam bahwa teknologi blockchain bukanlah obat mujarab (panacea) untuk segala masalah "
            "pencatatan data. Keputusan rekayasa yang matang terletak pada ketajaman membedakan kebutuhan integritas blockchain "
            "(ketahanan insider threat) versus basis data konvensional (transaksional tinggi). Pelajaran paling berharga adalah konfirmasi "
            "empiris bahwa gas eksekusi murni konstan O(1) (164.911 gas) dan selisih 84 gas deterministik dari calldata EIP-2028 [11], "
            "memperkuat pemahaman atas arsitektur internal EVM berbasis bukti nyata."
        )

    # Chapter 5.1 Intro
    if "Berdasarkan hasil perancangan arsitektur, implementasi smart contract, pengujian fungsional" in p.text or "Berdasarkan hasil perancangan arsitektur, implementasi smart contract, pengujian unit test" in p.text:
        p.text = "Berdasarkan hasil perancangan arsitektur, implementasi smart contract, pengujian fungsional, dan evaluasi performa, diperoleh lima kesimpulan utama:"

    # Chapter 5.1 Kesimpulan 1
    if "• 1. Kelayakan Arsitektur:" in p.text:
        p.text = (
            "• 1. Kelayakan Arsitektur: Sistem verifikasi berbasis blockchain konsorsium (Hyperledger Besu/EVM) dan pohon Merkle "
            "pada lingkungan uji ini menunjukkan kelayakan teknis, keandalan kontrol akses, dan skalabilitas tinggi dalam "
            "meniadakan risiko pemalsuan dokumen serta ketergantungan server terpusat."
        )

    # Chapter 5.1 Kesimpulan 2 (UU PDP separate articles)
    if "• 2. Perlindungan Privasi:" in p.text:
        p.text = (
            "• 2. Perlindungan Privasi: Arsitektur Zero-PII On-Chain yang memisahkan dokumen lengkap off-chain bergaram (256-bit "
            "salt) dari Merkle root on-chain [4] dirancang selaras dengan prinsip pelindungan data pribadi UU No. 27 Tahun 2022 (UU "
            "PDP) [10], meliputi perlindungan klasifikasi data pribadi umum (Pasal 4 ayat (3)), pemenuhan hak penghapusan data "
            "pribadi (Pasal 8 dan Pasal 43 ayat (1)), serta hak penarikan persetujuan (Pasal 9), sekaligus mengeliminasi ancaman "
            "serangan kamus terhadap atribut bernilai rendah."
        )

    # Chapter 5.1 Kesimpulan 3
    if "• 3. Skalabilitas O(1):" in p.text:
        p.text = (
            "• 3. Skalabilitas O(1): Efisiensi gas eksekusi penjangkaran batch wisuda terbukti konstan secara eksak sebesar "
            "164.911 gas (deviasi 0,0000%, true O(1)) untuk N=6, 100, hingga 1.000 lulusan. Selisih 84 gas pada tanda "
            "terima transaksi (receipt) dapat dijelaskan secara terukur bersumber dari penambahan panjang parameter "
            "calldata ABI transaksi sesuai aturan EIP-2028 [11]."
        )

    # Chapter 5.1 Kesimpulan 4
    if "• 4. Akses Verifikasi Gratis:" in p.text:
        p.text = (
            "• 4. Akses Verifikasi Gratis: Layanan verifikasi dokumen oleh masyarakat dan dunia industri dapat diakses secara instan "
            "(< 1 ms latensi kueri) dan bebas biaya (0 gas) melalui antarmuka JSON-RPC eth_call [2], dengan ukuran bukti pembuktian "
            "logaritmik O(log N) [4] yang sangat hemat bandwidth."
        )

    # Chapter 5.1 Kesimpulan 5
    if "• 5. Keamanan & Mitigasi:" in p.text:
        p.text = (
            "• 5. Keamanan & Mitigasi: Pada lingkungan uji ini, mekanisme keamanan smart contract menunjukkan "
            "ketahanan melalui penolakan seketika atas manipulasi klaim data (tamper-evident), perlindungan hak akses "
            "berbasis peran (RBAC), serta keandalan sirkuit pemutus darurat (circuit breaker) dalam merespons insiden "
            "keamanan dalam 24 jam pertama."
        )

    # Chapter 5.2 Saran Tindak Lanjut
    if "• 1. Implementasi Konsorsium:" in p.text:
        p.text = (
            "• 1. Implementasi Konsorsium: Mendorong pembentukan konsorsium resmi perguruan tinggi/LLDIKTI untuk uji coba "
            "rintisan multi-node nyata menggunakan Hyperledger Besu multi-validator."
        )
    if "• 2. Pengembangan Antarmuka:" in p.text:
        p.text = (
            "• 2. Pengembangan Antarmuka: Mengembangkan aplikasi dompet kredensial (credential wallet) mobile ramah "
            "pengguna berbasis standar W3C VC dan Open Badges v3.0."
        )
    if "• 3. Riset Lanjutan ZKP:" in p.text:
        p.text = (
            "• 3. Riset Lanjutan ZKP: Mengintegrasikan teknologi Selective Disclosure berbasis Zero-Knowledge Proofs (ZKP) "
            "agar mahasiswa dapat membuktikan kualifikasi tertentu (misal syarat akreditasi prodi atau batas minimum IPK) "
            "tanpa harus memperlihatkan seluruh riwayat transkrip nilai."
        )

    # Reference [10] in bibliography: set access date to 5 Oktober 2026
    if p.text.strip().startswith("[10]"):
        replace_in_p(p, "[Diakses: 4 Oktober 2026]", "[Diakses: 5 Oktober 2026]")

# =============================================================================
# 10. INSERT CONCISE CRITERIA SUMMARY IN SUBCHAPTER 4.6 (MAX 4 LINES)
# =============================================================================
for i, p in enumerate(doc.paragraphs):
    if "Kriteria penerimaan (acceptance criteria) disusun dengan merujuk" in p.text:
        p_crit = doc.add_paragraph()
        p_crit.text = (
            "Ringkasan ambang batas penerimaan luaran OBE (rincian lengkap disajikan pada Lampiran D):\n"
            "• Penjangkaran & Verifikasi (T-01, T-02): status=1, exists==true; isValid==true, status 'VALID', gas verifikator = 0 (via eth_call).\n"
            "• Integritas & Hak Akses (T-03, T-04, T-06): isValid==false ('INVALID_PROOF_OR_TAMPERED'); revert AccessControlUnauthorizedAccount.\n"
            "• Jeda & Pencabutan (T-05, T-07): Status 'CREDENTIAL_REVOKED' (O(1)); revert EnforcedPause saat jeda darurat aktif.\n"
            "• Skalabilitas & Anggaran (T-08): Deviasi gas murni <= 1,0%; bukti <= ceil(log2 N); pureGas(k) <= 10.000 + 350*k gas; bangun pohon < 1.000 ms."
        )
        p_crit.paragraph_format.space_before = Pt(2)
        p_crit.paragraph_format.space_after = Pt(2)
        p_crit.paragraph_format.line_spacing = 1.05
        for r in p_crit.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        p._element.addnext(p_crit._element)
        break

# =============================================================================
# 11. INSERT B-06 SUMMARY IN SUBCHAPTER 4.7 (MAX 3 SENTENCES + LAMPIRAN F)
# =============================================================================
for i, p in enumerate(doc.paragraphs):
    if "• 7. Efek Pembulatan Calldata" in p.text:
        p_b06 = doc.add_paragraph()
        p_b06.text = (
            "• 8. Analisis Selisih Gas Pencabutan B-06: Pada pengujian prototipe, revokeCredential mengonsumsi 56.359 gas "
            "pada demo [B-06] dan berkisar 55.297 s.d. 55.333 gas pada benchmark. Pengujian isolasi variabel mengonfirmasi "
            "selisih ini bersumber dari perbedaan calldata EIP-2028 serta lompatan 262 gas pada batas 32 byte data log event. "
            "Rincian tabel Run A s.d. E dan log pengujian disajikan pada Lampiran F."
        )
        p_b06.paragraph_format.space_before = Pt(1)
        p_b06.paragraph_format.space_after = Pt(1)
        p_b06.paragraph_format.line_spacing = 1.05
        for r in p_b06.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        p._element.addnext(p_b06._element)
        break

# =============================================================================
# 12. ADD PERMENDIKBUDRISTEK NO. 6/2022 TO DAFTAR PUSTAKA IF NOT PRESENT
# =============================================================================
p_ref12_idx = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("[12]"):
        p_ref12_idx = i
        break

if p_ref12_idx != -1:
    has_13 = any(p.text.strip().startswith("[13]") for p in doc.paragraphs)
    if not has_13:
        p13 = doc.add_paragraph()
        p13.text = (
            '[13] Kementerian Pendidikan, Kebudayaan, Riset, dan Teknologi Republik Indonesia, '
            '"Peraturan Menteri Pendidikan, Kebudayaan, Riset, dan Teknologi Republik Indonesia Nomor 6 Tahun 2022 '
            'tentang Ijazah, Sertifikat Kompetensi, Sertifikat Profesi, Gelar, dan Kesetaraan Ijazah Perguruan Tinggi Negara Lain," '
            'Berita Negara Republik Indonesia Tahun 2022 Nomor 167, Jakarta, Feb. 2022. [Daring]. Tersedia: '
            'https://peraturan.go.id/id/permendikbudristek-no-6-tahun-2022. [Diakses: 4 Oktober 2026].'
        )
        p13.paragraph_format.space_before = Pt(3)
        p13.paragraph_format.space_after = Pt(3)
        for r in p13.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
        doc.paragraphs[p_ref12_idx]._element.addnext(p13._element)

# =============================================================================
# 13. UPDATE LAMPIRAN A (PENGUNGKAPAN AI TRANSPARAN)
# =============================================================================
for i, p in enumerate(doc.paragraphs):
    if "Lampiran A: Pengungkapan Penggunaan Alat Bantu AI" in p.text:
        # Find the bullets under Lampiran A
        for j in range(i + 1, min(i + 10, len(doc.paragraphs))):
            pj = doc.paragraphs[j]
            if "• 2. Google Antigravity Agent:" in pj.text:
                # Add bullet 3 with honest v7-v7.2 disclosure
                p_ai = doc.add_paragraph()
                p_ai.text = (
                    "• 3. Siklus Koreksi & Verifikasi Laporan (v7 s.d. v7.2): Tahap koreksi v7 sampai v7.2 "
                    "dikerjakan dengan agent (Google Antigravity) atas arahan saya, dan peninjauan serta penyusunan "
                    "instruksi koreksi dibantu Claude (asisten AI Anthropic); kode kontrak, hasil uji, dan angka "
                    "pengukuran tidak diubah."
                )
                p_ai.paragraph_format.space_before = Pt(2)
                p_ai.paragraph_format.space_after = Pt(2)
                for r in p_ai.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(8.5)
                pj._element.addnext(p_ai._element)
            if "• 3. Tanggung Jawab Penulis:" in pj.text:
                pj.text = pj.text.replace("• 3. Tanggung Jawab Penulis:", "• 4. Tanggung Jawab Penulis:")
                for r in pj.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(8.5)
        break

# =============================================================================
# 14. INSERT NUMBERED CAPTIONS FOR BODY TABLES 1 THROUGH 10
# =============================================================================
body_captions = [
    ("Komponen", "Tabel 1. Komponen Lingkungan Pengujian dan Reproduksibilitas Sistem"),
    ("Tahap", "Tabel 2. Catatan Pelaksanaan Mini-Project Perancangan Sistem Blockchain"),
    ("Aspek Evaluasi", "Tabel 3. Perbandingan Karakteristik Basis Data Terpusat vs Blockchain Konsorsium"),
    ("Komponen Data / Fungsi", "Tabel 4. Partisi Komponen Data On-Chain vs Off-Chain dan Kepatuhan Privasi"),
    ("Record / Daun", "Tabel 5. Struktur Data Record Daun Merkle Kredensial Akademik Uji Coba"),
    ("Mekanisme Konsensus", "Tabel 6. Analisis Komparatif Mekanisme Konsensus Terdistribusi"),
    ("Field Transaksi EVM", "Tabel 7. Parameter Field Transaksi Penjangkaran EVM pada Node Lokal"),
    ("Komponen Gas Transaksi issueBatch", "Tabel 8. Rincian Komponen Konsumsi Gas Transaksi issueBatch (N = 6, 100, 1.000)"),
    ("Kategori Ancaman", "Tabel 9. Matriks Analisis Ancaman STRIDE dan Kontrol Mitigasi Sistem"),
    ("ID Tes", "Tabel 10. Ringkasan Hasil Eksekusi Pengujian Unit Test Otomatis (T-01 s.d. T-08)")
]

for identifier, cap_text in body_captions:
    target_tbl = None
    for t in doc.tables:
        if identifier in t.rows[0].cells[0].text:
            target_tbl = t
            break
    if target_tbl is not None:
        p_cap = doc.add_paragraph()
        p_cap.text = cap_text
        p_cap.paragraph_format.keep_with_next = True
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        for r in p_cap.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        target_tbl._element.addprevious(p_cap._element)

# =============================================================================
# 15. SCALE DIAGRAM 2 (80%)
# =============================================================================
scale = 0.80
for p in doc.paragraphs:
    for r in p.runs:
        extents = r._r.xpath(".//wp:extent")
        for ext in extents:
            cx = int(ext.attrib.get("cx", 0))
            cy = int(ext.attrib.get("cy", 0))
            if cy > 3000000:
                ext.attrib["cx"] = str(int(cx * scale))
                ext.attrib["cy"] = str(int(cy * scale))

# =============================================================================
# 16. RAISE FONT SIZE FOR LAMPIRAN D AND E TO MINIMAL 9 PT
# =============================================================================
for t in doc.tables:
    first_cell = t.rows[0].cells[0].text.strip()
    if "Kelompok Uji" in first_cell or "Parameter Pengujian" in first_cell:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9.0)
                        if "30M gas" in run.text:
                            run.text = run.text.replace("30M gas", "60M gas node lokal / 30M gas contoh hipotetis konsorsium")

# =============================================================================
# 17. CREATE LAMPIRAN F (EKSPERIMEN B-06)
# =============================================================================
p_lamp_f_heading = doc.add_paragraph()
p_lamp_f_heading.text = "Lampiran F: Eksperimen Isolasi Variabel Selisih Gas revokeCredential (B-06)"
p_lamp_f_heading.paragraph_format.space_before = Pt(16)
p_lamp_f_heading.paragraph_format.space_after = Pt(4)
p_lamp_f_heading.paragraph_format.keep_with_next = True
for r in p_lamp_f_heading.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

p_lamp_f_intro = doc.add_paragraph()
p_lamp_f_intro.text = (
    "Eksperimen isolasi variabel tunggal (exp_b06_v2.js) dirancang untuk memverifikasi akar penyebab matematis "
    "perbedaan konsumsi gas fungsi revokeCredential antara skrip demo (56.359 gas pada Bukti B-06) dan hasil benchmark "
    "(55.297 s.d. 55.333 gas). Seluruh run dieksekusi pada smart contract yang sama (AcademicCredentialRegistry.sol), "
    "pemanggil yang sama (admin akun penerbit berwenang), dan status state awal yang identik melalui mekanisme EVM snapshot/revert. "
    "Panjang parameter diukur secara eksak menggunakan Buffer.byteLength() dan komponen calldata dihitung sesuai aturan EIP-2028."
)
p_lamp_f_intro.paragraph_format.space_before = Pt(2)
p_lamp_f_intro.paragraph_format.space_after = Pt(4)
for r in p_lamp_f_intro.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(9.0)

# Table Run A to E
table_f_data = [
    ["Skenario Uji (Run)", "Parameter batchId", "Parameter reason", "Calldata Gas", "Gas Receipt", "Selisih vs Run A", "Prediksi Δ Cd", "SISA (Terukur-ΔCd)"],
    ["Run A (Identik Demo)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "Putusan Sidang Komite Etik... (60 B)", "2.212 gas", "56.359 gas", "0 (Basis)", "0 gas", "0 gas"],
    ["Run B (Ubah batchId)", "BATCH-6-ITER-1 (32 B, 14 non-0)", "Putusan Sidang Komite Etik... (60 B)", "2.128 gas", "56.275 gas", "-84 gas", "-84 gas", "0 gas"],
    ["Run C (Ubah reason)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "Sidang Etik (11 B)", "1.496 gas", "55.381 gas", "-978 gas", "-716 gas", "-262 gas"],
    ["Run D (Reason 11 B)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "Sidang Etik (11 B UTF-8)", "1.496 gas", "55.381 gas", "-978 gas", "-716 gas", "-262 gas"],
    ["Run D (Reason 31 B)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "31 byte string UTF-8", "1.736 gas", "55.621 gas", "-738 gas", "-476 gas", "-262 gas"],
    ["Run D (Reason 32 B)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "32 byte string UTF-8", "1.748 gas", "55.633 gas", "-726 gas", "-464 gas", "-262 gas"],
    ["Run D (Reason 33 B)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "33 byte string UTF-8", "1.888 gas", "56.035 gas", "-324 gas", "-324 gas", "0 gas"],
    ["Run D (Reason 60 B)", "WISUDA-2026-PERIODE-1 (32 B, 21 non-0)", "60 byte string UTF-8 (Demo)", "2.212 gas", "56.359 gas", "0 (Basis)", "0 gas", "0 gas"],
    ["Run E (Bench N=6)", "BATCH-6-ITER-1 (32 B, 14 non-0)", "Sidang Etik (11 B UTF-8)", "1.412 gas", "55.297 gas", "-1.062 gas", "-800 gas", "-262 gas"],
    ["Run E (Bench N=100)", "BATCH-100-ITER-1 (32 B, 16 non-0)", "Sidang Etik (11 B UTF-8)", "1.436 gas", "55.321 gas", "-1.038 gas", "-776 gas", "-262 gas"],
    ["Run E (Bench N=1000)", "BATCH-1000-ITER-1 (32 B, 17 non-0)", "Sidang Etik (11 B UTF-8)", "1.448 gas", "55.333 gas", "-1.026 gas", "-764 gas", "-262 gas"]
]

tbl_f = doc.add_table(rows=len(table_f_data), cols=len(table_f_data[0]))
tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER

for row_idx, row in enumerate(table_f_data):
    for col_idx, text in enumerate(row):
        cell = tbl_f.cell(row_idx, col_idx)
        cell.text = text
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(8.5 if row_idx > 0 else 9.0)
                if row_idx == 0:
                    r.font.bold = True

p_lamp_f_analysis = doc.add_paragraph()
p_lamp_f_analysis.text = (
    "Analisis Hasil Eksperimen dan Temuan Matematis Terukur:\n"
    "1. Resolusi Panjang String: String reason demo terkonfirmasi berukuran tepat 60 karakter / 60 byte UTF-8 "
    "(perbedaan dengan angka 61 pada laporan sebelumnya terselesaikan sebagai salah hitung manual).\n"
    "2. Variasi batchId (Run B): Perubahan batchId dari format demo ke benchmark menghasilkan selisih tepat -84 gas, "
    "yang 100% cocok dengan penurunan 7 byte non-nol calldata EIP-2028 (7 * 12 gas = 84 gas) dengan SISA = 0 gas.\n"
    "3. Lompatan Batas 32-Byte Event Log Data (Run D): Pada pengujian Run D, saat panjang string reason dinaikkan dari "
    "32 byte menjadi 33 byte, terjadi lonjakan tepat 262 gas pada komponen non-calldata (nilai SISA berubah dari -262 menjadi 0). "
    "Lompatan 262 gas ini terkonfirmasi terdiri atas 256 gas (satu word 32-byte tambahan pada data event CredentialRevoked "
    "sesuai formula EVM log data: 8 gas/byte * 32 byte = 256 gas) + 6 gas (belum dirinci).\n"
    "4. SISA Non-Nol: Komponen di luar biaya calldata EIP-2028 dan alokasi data event log dinyatakan belum dijelaskan."
)
p_lamp_f_analysis.paragraph_format.space_before = Pt(6)
p_lamp_f_analysis.paragraph_format.space_after = Pt(4)
for r in p_lamp_f_analysis.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(9.0)

p_lamp_f_log_hdr = doc.add_paragraph()
p_lamp_f_log_hdr.text = "Keluaran Mentah Eksekusi Terminal (uts-blockchain-submission/tools/exp_b06_output.txt):"
p_lamp_f_log_hdr.paragraph_format.space_before = Pt(6)
p_lamp_f_log_hdr.paragraph_format.space_after = Pt(2)
for r in p_lamp_f_log_hdr.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(9.0)
    r.font.bold = True

# Read raw log
with open("uts-blockchain-submission/tools/exp_b06_output.txt", "r", encoding="utf-8") as fp:
    raw_log = fp.read()

p_lamp_f_log = doc.add_paragraph()
p_lamp_f_log.text = raw_log
p_lamp_f_log.paragraph_format.space_before = Pt(2)
p_lamp_f_log.paragraph_format.space_after = Pt(4)
for r in p_lamp_f_log.runs:
    r.font.name = 'Courier New'
    r.font.size = Pt(8.0)

# =============================================================================
# 18. REMOVE REDUNDANT PARAGRAPHS & TIGHTEN HEADING SPACING
# =============================================================================
to_remove = []
for p in doc.paragraphs:
    if "Seluruh rangkaian unit test otomatis (T-01 s.d. T-08)" in p.text:
        to_remove.append(p)
    if p.text.strip() == "" and "drawing" not in p._p.xml:
        to_remove.append(p)

for p in to_remove:
    p._element.getparent().remove(p._element)

for p in doc.paragraphs:
    if p.paragraph_format.space_before and p.paragraph_format.space_before.pt > 4:
        p.paragraph_format.space_before = Pt(3)
    if p.paragraph_format.space_after and p.paragraph_format.space_after.pt > 3:
        p.paragraph_format.space_after = Pt(2)

# =============================================================================
# 19. SAVE DOCX AND CONVERT TO PDF
# =============================================================================
doc.save(OUT_DOCX)
print(f"v7.2 DOCX saved successfully to {OUT_DOCX}")

cmd = ["soffice", "--headless", "--convert-to", "pdf", OUT_DOCX, "--outdir", "v7_workspace"]
subprocess.run(cmd, check=True)
print(f"v7.2 PDF compiled successfully to {OUT_PDF}")

# Apply Metadata via Exiftool
exif_cmd = [
    "exiftool",
    "-Author=Fajar Geran Arifin",
    "-Title=UTS Blockchain - Rancangan Sistem Verifikasi Kredensial Akademik",
    "-overwrite_original",
    OUT_PDF
]
subprocess.run(exif_cmd, check=True)
print("PDF metadata (Author & Title) updated successfully via exiftool.")
